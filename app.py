"""ATS Resume Checker
Upload a resume (PDF / DOCX / TXT) and get an ATS-style score plus concrete
improvements. UI: Streamlit. AI: Groq.

Run locally:  streamlit run app.py
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import re

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
DEFAULT_MODELS = [
    "llama-3.3-70b-versatile",
    "llama-3.1-8b-instant",
    "openai/gpt-oss-120b",
]
CUSTOM_MODEL_LABEL = "Custom..."

# Kept modest so requests fit inside Groq's free-tier tokens-per-minute limits.
MAX_RESUME_CHARS = 9000
MAX_JD_CHARS = 4000
MIN_TEXT_CHARS = 150

# Final score = RULE_WEIGHT * format score + (1 - RULE_WEIGHT) * AI content score
RULE_WEIGHT = 0.4

SECTION_KEYWORDS = {
    "Experience": ["experience", "employment", "work history"],
    "Education": ["education", "academic", "qualification"],
    "Skills": ["skills", "technologies", "technical proficiencies", "competencies"],
    "Summary / Objective": ["summary", "objective", "about me", "professional profile"],
    "Projects / Certifications": [
        "projects",
        "certification",
        "certificates",
        "achievements",
        "publications",
    ],
}
SECTION_POINTS = {
    "Experience": 15,
    "Education": 10,
    "Skills": 15,
    "Summary / Objective": 5,
    "Projects / Certifications": 5,
}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"\+?\d[\d\s().-]{8,}\d")
LINK_RE = re.compile(r"(linkedin\.com|github\.com)", re.I)
BULLET_RE = re.compile(r"^\s*[-\u2022\u25cf\u25aa\u25e6\u2013*]\s+\S", re.M)
METRIC_RE = re.compile(
    r"\d[\d,.]*\s?(?:%|\+|x\b|k\b|m\b|million|thousand|users|customers|clients|"
    r"projects|hours|days|ms\b|seconds)|[$\u00a3\u20ac]\s?\d",
    re.I,
)

SYSTEM_PROMPT = (
    "You are an expert technical recruiter and ATS (Applicant Tracking System) "
    "specialist. You respond with a single valid JSON object and nothing else."
)


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(file_name: str, data: bytes) -> str:
    """Return the plain text of a PDF, DOCX or TXT resume."""
    ext = os.path.splitext(file_name.lower())[1]

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ValueError("This PDF is password-protected.") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
        return clean_text("\n".join(pages))

    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return clean_text("\n".join(parts))

    if ext == ".txt":
        return clean_text(data.decode("utf-8", errors="replace"))

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------- #
# Rule-based checks (format & structure, 0-100)
# --------------------------------------------------------------------------- #
def _has_phone(text: str) -> bool:
    for match in PHONE_RE.finditer(text):
        groups = re.findall(r"\d+", match.group())
        if all(len(g) == 4 and g[:2] in ("19", "20") for g in groups):
            continue  # just a run of years, e.g. "2019 - 2022 - 2024"
        if 10 <= len("".join(groups)) <= 15:
            return True
    return False


def _find_sections(text: str) -> set[str]:
    found = set()
    for line in text.splitlines():
        line = line.strip().strip(":").lower()
        if not line or len(line.split()) > 4 or line.endswith("."):
            continue
        for section, keywords in SECTION_KEYWORDS.items():
            if any(k in line for k in keywords):
                found.add(section)
    return found


def rule_based_checks(text: str) -> dict:
    checks: list[dict] = []

    def add(name: str, points: int, max_points: int, tip: str) -> None:
        checks.append(
            {
                "name": name,
                "points": points,
                "max": max_points,
                "passed": points == max_points,
                "tip": tip,
            }
        )

    add(
        "Email address",
        10 if EMAIL_RE.search(text) else 0,
        10,
        "Add a professional email address in plain text at the top.",
    )
    add(
        "Phone number",
        10 if _has_phone(text) else 0,
        10,
        "Add a phone number (include the country code if applying abroad).",
    )
    add(
        "LinkedIn / GitHub link",
        5 if LINK_RE.search(text) else 0,
        5,
        "Add your LinkedIn and/or GitHub URL.",
    )

    sections = _find_sections(text)
    for section, pts in SECTION_POINTS.items():
        add(
            f"Section: {section}",
            pts if section in sections else 0,
            pts,
            f"Add a clearly labelled '{section.split(' / ')[0]}' heading so ATS can find it.",
        )

    words = len(text.split())
    if 300 <= words <= 900:
        length_pts = 10
    elif 200 <= words < 300 or 900 < words <= 1300:
        length_pts = 6
    else:
        length_pts = 2
    add(
        f"Length ({words} words)",
        length_pts,
        10,
        "Aim for roughly 300-900 words (1-2 pages).",
    )

    bullets = len(BULLET_RE.findall(text))
    bullet_pts = 5 if bullets >= 5 else 2 if bullets >= 1 else 0
    add(
        f"Bullet points ({bullets} found)",
        bullet_pts,
        5,
        "Use bullet points for responsibilities and achievements.",
    )

    metrics = len(METRIC_RE.findall(text))
    metric_pts = 10 if metrics >= 3 else 5 if metrics >= 1 else 0
    add(
        f"Quantified achievements ({metrics} found)",
        metric_pts,
        10,
        "Add numbers: percentages, money, users, time saved, team size.",
    )

    return {"score": sum(c["points"] for c in checks), "checks": checks}


# --------------------------------------------------------------------------- #
# Groq / LLM analysis
# --------------------------------------------------------------------------- #
def build_user_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"\nJOB DESCRIPTION (score job_match_score against this):\n"
        f"\"\"\"\n{job_description}\n\"\"\"\n"
        if job_description
        else "\nNo job description was provided, so set job_match_score to null.\n"
    )
    return f"""Evaluate the resume below for ATS compatibility and content quality.
{jd_block}
RESUME:
\"\"\"
{resume_text}
\"\"\"

Rules:
- Be strict and realistic. Most real resumes score between 55 and 80. Do not inflate scores.
- Base every judgement only on the resume text. Never invent experience, employers or numbers.
- In bullet_rewrites, improve wording only: start with a strong action verb and, where a metric
  is missing, use a placeholder such as [X%] or [N users] instead of making one up.

Return ONE JSON object with exactly these keys:
{{
  "overall_score": integer 0-100 (content quality and ATS-friendliness of wording),
  "keyword_score": integer 0-100 (relevant skills / industry keywords),
  "impact_score": integer 0-100 (achievements, results, metrics),
  "clarity_score": integer 0-100 (readability, structure, concise language),
  "job_match_score": integer 0-100 or null,
  "summary": "2-3 sentence overall assessment",
  "strengths": ["up to 5 short strings"],
  "improvements": [
    {{"priority": "High|Medium|Low", "area": "short label", "issue": "what is wrong",
      "suggestion": "specific fix"}}
  ],
  "missing_keywords": ["up to 15 keywords or skills worth adding if truthful"],
  "bullet_rewrites": [{{"original": "exact weak line from the resume", "improved": "better version"}}]
}}
Give at most 8 improvements and at most 5 bullet_rewrites."""


def parse_json_response(raw: str) -> dict:
    """Parse the model output into a dict, tolerating code fences and extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("The AI response was not valid JSON.") from None
        try:
            data = json.loads(raw[start : end + 1])
        except json.JSONDecodeError as exc:
            raise ValueError("The AI response was not valid JSON.") from exc
    if not isinstance(data, dict):
        raise ValueError("The AI response had an unexpected format.")
    return data


def _clamp(value, default=None):
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value, limit: int) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if isinstance(v, (str, int, float)) and str(v).strip()][
        :limit
    ]


def normalize_result(data: dict, has_jd: bool) -> dict:
    """Validate / clamp the model output so the UI never breaks."""
    keyword = _clamp(data.get("keyword_score"))
    impact = _clamp(data.get("impact_score"))
    clarity = _clamp(data.get("clarity_score"))

    overall = _clamp(data.get("overall_score"))
    if overall is None:
        subs = [s for s in (keyword, impact, clarity) if s is not None]
        if not subs:
            raise ValueError("The AI response did not include any scores.")
        overall = round(sum(subs) / len(subs))

    priority_order = {"High": 0, "Medium": 1, "Low": 2}
    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in priority_order:
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "area": str(item.get("area", "General")).strip() or "General",
                "issue": str(item.get("issue", "")).strip(),
                "suggestion": str(item.get("suggestion", "")).strip(),
            }
        )
    improvements.sort(key=lambda i: priority_order[i["priority"]])

    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]).strip(), "improved": str(item["improved"]).strip()}
            )

    return {
        "overall_score": overall,
        "keyword_score": keyword,
        "impact_score": impact,
        "clarity_score": clarity,
        "job_match_score": _clamp(data.get("job_match_score")) if has_jd else None,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths"), 5),
        "improvements": improvements[:8],
        "missing_keywords": _str_list(data.get("missing_keywords"), 15),
        "bullet_rewrites": rewrites[:5],
    }


def analyze_with_groq(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    client = Groq(api_key=api_key)
    kwargs = {
        "model": model,
        "temperature": 0.2,
        "max_tokens": 2000,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": build_user_prompt(
                    resume_text[:MAX_RESUME_CHARS], job_description[:MAX_JD_CHARS]
                ),
            },
        ],
    }
    try:
        response = client.chat.completions.create(
            **kwargs, response_format={"type": "json_object"}
        )
    except Exception as exc:  # noqa: BLE001
        # Some models / transient JSON-mode failures: retry once in plain mode.
        message = str(exc).lower()
        if "response_format" in message or "json" in message:
            response = client.chat.completions.create(**kwargs)
        else:
            raise
    raw = response.choices[0].message.content
    return normalize_result(parse_json_response(raw), has_jd=bool(job_description.strip()))


def combine_scores(rule_score: int, ai_score: int) -> int:
    return round(RULE_WEIGHT * rule_score + (1 - RULE_WEIGHT) * ai_score)


def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent - likely to pass most ATS filters"
    if score >= 70:
        return "Good - a few improvements will make it stronger"
    if score >= 50:
        return "Average - needs noticeable work"
    return "Needs a major revision"


def friendly_error(exc: Exception) -> str:
    name = type(exc).__name__
    if name == "AuthenticationError":
        return "Groq rejected the API key. Please check it and try again."
    if name == "RateLimitError":
        return "Groq rate limit reached. Wait a minute, or choose a smaller model, then retry."
    if name in ("APIConnectionError", "APITimeoutError"):
        return "Could not reach Groq. Check your internet connection and retry."
    if name in ("NotFoundError", "BadRequestError") and "model" in str(exc).lower():
        return "That model is not available on your Groq account. Pick another model."
    return f"{name}: {exc}"


def build_report_md(file_name: str, res: dict) -> str:
    ai = res["ai"]
    lines = [
        f"# ATS Resume Report - {file_name}",
        "",
        f"**Overall ATS score:** {res['final_score']}/100 ({score_label(res['final_score'])})",
        f"- Format & structure: {res['rule']['score']}/100",
        f"- Content quality (AI): {ai['overall_score']}/100",
    ]
    if ai["job_match_score"] is not None:
        lines.append(f"- Job match: {ai['job_match_score']}/100")
    lines += ["", "## Summary", ai["summary"] or "-", "", "## Improvements"]
    for i in ai["improvements"]:
        lines.append(f"- **[{i['priority']}] {i['area']}** - {i['issue']} Fix: {i['suggestion']}")
    lines += ["", "## Missing keywords", ", ".join(ai["missing_keywords"]) or "-", ""]
    lines.append("## Bullet rewrites")
    for r in ai["bullet_rewrites"]:
        lines += [f"- Before: {r['original']}", f"  After: {r['improved']}"]
    lines += ["", "## Format checks"]
    for c in res["rule"]["checks"]:
        mark = "PASS" if c["passed"] else "FIX"
        lines.append(f"- [{mark}] {c['name']} ({c['points']}/{c['max']})")
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Streamlit UI
# --------------------------------------------------------------------------- #
def get_default_api_key() -> str:
    try:
        if "GROQ_API_KEY" in st.secrets:
            return str(st.secrets["GROQ_API_KEY"])
    except Exception:  # noqa: BLE001  (no secrets file configured)
        pass
    return os.getenv("GROQ_API_KEY", "")


def render_results(res: dict, file_name: str) -> None:
    ai = res["ai"]
    final = res["final_score"]

    st.divider()
    st.header("Results")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Overall ATS score", f"{final}/100")
    c2.metric("Format & structure", f"{res['rule']['score']}/100")
    c3.metric("Content quality (AI)", f"{ai['overall_score']}/100")
    c4.metric(
        "Job match",
        f"{ai['job_match_score']}/100" if ai["job_match_score"] is not None else "n/a",
    )
    st.progress(final / 100)
    st.caption(score_label(final))

    sub = st.columns(3)
    sub[0].metric("Keywords", f"{ai['keyword_score']}/100" if ai["keyword_score"] is not None else "n/a")
    sub[1].metric("Impact", f"{ai['impact_score']}/100" if ai["impact_score"] is not None else "n/a")
    sub[2].metric("Clarity", f"{ai['clarity_score']}/100" if ai["clarity_score"] is not None else "n/a")

    tabs = st.tabs(
        ["Improvements", "Summary & strengths", "Keywords", "Bullet rewrites", "Format checks", "Extracted text"]
    )

    with tabs[0]:
        if not ai["improvements"]:
            st.write("No improvements returned.")
        icons = {"High": "\U0001f534", "Medium": "\U0001f7e0", "Low": "\U0001f7e2"}
        for item in ai["improvements"]:
            with st.expander(f"{icons[item['priority']]} {item['priority']} - {item['area']}"):
                st.markdown(f"**Issue:** {item['issue']}")
                st.markdown(f"**Fix:** {item['suggestion']}")

    with tabs[1]:
        st.write(ai["summary"] or "No summary returned.")
        if ai["strengths"]:
            st.subheader("Strengths")
            for s in ai["strengths"]:
                st.markdown(f"- {s}")

    with tabs[2]:
        if ai["missing_keywords"]:
            st.write("Consider adding these **only if they are true for you**:")
            st.write(", ".join(f"`{k}`" for k in ai["missing_keywords"]))
        else:
            st.write("No missing keywords reported.")

    with tabs[3]:
        if not ai["bullet_rewrites"]:
            st.write("No rewrites returned.")
        for r in ai["bullet_rewrites"]:
            st.markdown(f"**Before:** {r['original']}")
            st.markdown(f"**After:** {r['improved']}")
            st.divider()

    with tabs[4]:
        for c in res["rule"]["checks"]:
            if c["passed"]:
                st.markdown(f"\u2705 **{c['name']}** ({c['points']}/{c['max']})")
            else:
                st.markdown(f"\u274c **{c['name']}** ({c['points']}/{c['max']}) - {c['tip']}")

    with tabs[5]:
        st.caption("This is what an ATS 'sees'. If it looks garbled, your layout may be hard to parse.")
        st.text_area("Extracted text", res["text"], height=300, label_visibility="collapsed")

    st.download_button(
        "Download report (.md)",
        data=build_report_md(file_name, res),
        file_name="ats_report.md",
        mime="text/markdown",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="\U0001f4c4", layout="wide")
    st.title("\U0001f4c4 ATS Resume Checker")
    st.write("Upload your resume to get an estimated ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = st.text_input(
            "Groq API key",
            type="password",
            value=get_default_api_key(),
            help="Get a free key at console.groq.com/keys",
        )
        choice = st.selectbox("Model", DEFAULT_MODELS + [CUSTOM_MODEL_LABEL])
        model = (
            st.text_input("Custom model ID", value=DEFAULT_MODELS[0])
            if choice == CUSTOM_MODEL_LABEL
            else choice
        )
        st.caption(
            "Your resume text is sent to Groq for analysis. Do not upload documents "
            "you are not comfortable sharing."
        )

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, for a tailored job-match score)",
        height=150,
        placeholder="Paste the job posting here...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        data = uploaded.getvalue()
        try:
            text = extract_text(uploaded.name, data)
        except Exception as exc:  # noqa: BLE001
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < MIN_TEXT_CHARS:
            st.error(
                "Almost no text could be extracted. If this is a scanned/image PDF, an ATS "
                "cannot read it either - export a text-based PDF or upload a DOCX."
            )
            return
        if not api_key.strip():
            st.error("Please enter your Groq API key in the sidebar.")
            return
        if len(text) > MAX_RESUME_CHARS:
            st.info(f"Long resume: only the first {MAX_RESUME_CHARS:,} characters were analyzed.")

        try:
            with st.spinner("Analyzing your resume..."):
                ai = analyze_with_groq(api_key.strip(), model.strip(), text, job_description)
        except Exception as exc:  # noqa: BLE001
            st.error(friendly_error(exc))
            return

        rule = rule_based_checks(text)
        st.session_state["result"] = {
            "file_hash": hashlib.sha256(data).hexdigest(),
            "file_name": uploaded.name,
            "text": text,
            "rule": rule,
            "ai": ai,
            "final_score": combine_scores(rule["score"], ai["overall_score"]),
        }

    result = st.session_state.get("result")
    if (
        result
        and uploaded is not None
        and hashlib.sha256(uploaded.getvalue()).hexdigest() == result["file_hash"]
    ):
        render_results(result, result["file_name"])
        st.caption(
            "No single universal 'ATS score' exists - every ATS works differently. "
            "Treat this as a practical estimate, not a guarantee."
        )


if __name__ == "__main__":
    main()
